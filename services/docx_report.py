# -*- coding: utf-8 -*-
"""报告 docx 转换（内容级格式：标题/表格/页码；audit-delivery-pipeline 管线口径，禁 Word COM）。"""
import os, re

def md_to_docx(md_path, docx_path=None):
    from docx import Document
    from docx.shared import Pt, Cm
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn

    docx_path = docx_path or os.path.splitext(md_path)[0] + ".docx"
    with open(md_path, encoding="utf-8") as f:
        lines = f.read().splitlines()

    doc = Document()
    for sec in doc.sections:  # 边距对齐管线规范：上下2.54 左右3.18
        sec.top_margin = sec.bottom_margin = Cm(2.54)
        sec.left_margin = sec.right_margin = Cm(3.18)

    def _cn(run, font="仿宋", size=12, bold=False):
        run.font.name = "Times New Roman"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), font)
        run.font.size = Pt(size)
        run.bold = bold

    def _strip(s):
        return s.replace("**", "")

    # 页脚页码：fldSimple PAGE/NUMPAGES
    footer = doc.sections[0].footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    from docx.oxml import OxmlElement
    for el, txt in (("PAGE", "第"), ("NUMPAGES", "页 共")):
        fld = OxmlElement("w:fldSimple")
        fld.set(qn("w:instr"), el)
        r = OxmlElement("w:r"); t = OxmlElement("w:t"); t.text = "1"
        r.append(t); fld.append(r)
        footer._p.append(fld)
        _cn(footer.add_run(txt), size=9)
    _cn(footer.add_run("页"), size=9)

    def add_table(rows):
        cells = [[c.strip() for c in r.strip().strip("|").split("|")] for r in rows]
        cells = [r for r in cells if not all(re.fullmatch(r":?-{2,}:?", c or "---") for c in r)]
        if not cells:
            return
        ncol = max(len(r) for r in cells)
        tb = doc.add_table(rows=len(cells), cols=ncol)
        tb.style = "Table Grid"
        for i, row in enumerate(cells):
            for j in range(ncol):
                p = tb.cell(i, j).paragraphs[0]
                run = p.add_run(_strip(row[j]) if j < len(row) else "")
                _cn(run, font="宋体", size=10.5, bold=(i == 0))

    i = 0
    while i < len(lines):
        ln = lines[i]
        if ln.startswith("|"):
            j = i
            while j < len(lines) and lines[j].startswith("|"):
                j += 1
            add_table(lines[i:j]); i = j
            continue
        if not ln.strip():
            i += 1; continue
        if ln.startswith("# "):
            p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            _cn(p.add_run(_strip(ln[2:])), font="黑体", size=14, bold=True)
        elif ln.startswith("## "):
            _cn(doc.add_paragraph().add_run(_strip(ln[3:])), font="黑体", size=12, bold=True)
        elif ln.startswith("### "):
            _cn(doc.add_paragraph().add_run(_strip(ln[4:])), font="黑体", size=12, bold=True)
        elif ln.startswith("- "):
            p = doc.add_paragraph(_strip(ln[2:]))
            _cn(p.runs[0])
            p.paragraph_format.first_line_indent = Pt(24)
        elif ln.strip() == "---":
            pass
        else:
            txt = _strip(ln)
            m = re.match(r"^\*\*(.+?)\*\*$", ln.strip())
            if m:
                txt = m.group(1)
            p = doc.add_paragraph(txt)
            if p.runs:
                _cn(p.runs[0])
            p.paragraph_format.first_line_indent = Pt(24)  # 首行缩进2字符
            p.paragraph_format.line_spacing = 1.5
        i += 1

    doc.save(docx_path)
    return docx_path

if __name__ == "__main__":
    import sys
    p = md_to_docx(sys.argv[1])
    from docx import Document
    d = Document(p)
    print("paras:", len(d.paragraphs), "tables:", len(d.tables), "->", p)
